from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.services.ingestion.normalizer import normalize_row
from app.services.ingestion.pdf import _header_map


def parse_excel(path: str | Path, fallback_term: str = "") -> list[dict]:
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    rows: list[dict] = []
    for sheet in wb.worksheets:
        header_idx = None
        mapping = None
        data = list(sheet.iter_rows(values_only=True))
        for idx, row in enumerate(data[:15]):
            mapping = _header_map([str(c) if c is not None else "" for c in row])
            if mapping:
                header_idx = idx
                break
        if mapping is None or header_idx is None:
            continue
        for row in data[header_idx + 1 :]:
            if not row:
                continue
            cells = [str(c).strip() if c is not None else "" for c in row]

            def get(field: str) -> str:
                idx = mapping.get(field)
                return cells[idx] if idx is not None and idx < len(cells) else ""

            if not get("name"):
                continue
            rows.append(
                normalize_row(
                    {
                        "name": get("name"),
                        "code": get("code"),
                        "credit": get("credit"),
                        "term": get("term") or fallback_term,
                        "status": get("status") or None,
                        "score": get("score") or None,
                        "nature": get("nature"),
                        "category": get("category"),
                    },
                    fallback_term,
                )
            )
    wb.close()
    return [r for r in rows if r["raw_name"] and r["raw_credit"] is not None]
